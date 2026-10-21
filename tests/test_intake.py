"""Real source extraction and conservative parser regression checks."""
from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from questionbank.intake import extract_source, parse_question_text, strip_personal_answer_text


PANDOC = shutil.which("pandoc")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class ExtractionTests(unittest.TestCase):
    def setUp(self) -> None:
        # Windows sandbox temp may be readable but deny creating child files.
        # Keep disposable fixtures under this project's writable test root.
        self.temp_base = Path(__file__).resolve().parents[1] / ".test-tmp"
        self.temp_base.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=self.temp_base)
        self.root = Path(self.temp.name)
        self.work = self.root / "derived"

    def tearDown(self) -> None:
        self.assertTrue(self.root.resolve().is_relative_to(self.temp_base.resolve()))
        self.temp.cleanup()

    def extract_unchanged(self, path: Path) -> dict:
        before = digest(path)
        result = extract_source(path, self.work, PANDOC)
        self.assertEqual(before, digest(path))
        self.assertEqual(result["source"]["sha256"], before)
        self.assertEqual(result["source"]["path"], str(path.resolve()))
        return result

    def test_utf8_bom_and_gb18030_text(self) -> None:
        utf8 = self.root / "bom.txt"
        utf8.write_bytes("1.【填空】数值？\n答案：1.5".encode("utf-8-sig"))
        extracted = self.extract_unchanged(utf8)
        self.assertTrue(extracted["text"].startswith("1."))
        self.assertFalse(extracted["needsVisualReview"])
        legacy = self.root / "legacy.txt"
        legacy.write_bytes("中文题库\n答案：两个词".encode("gb18030"))
        result = self.extract_unchanged(legacy)
        self.assertIn("中文题库", result["text"])
        self.assertTrue(any("GB18030" in note for note in result["notes"]))

    def test_json_is_read_verbatim_even_when_invalid(self) -> None:
        valid = self.root / "valid.json"
        valid.write_text(json.dumps({"title": "知识", "items": []}, ensure_ascii=False), encoding="utf-8")
        self.assertEqual(json.loads(self.extract_unchanged(valid)["text"])["title"], "知识")
        invalid = self.root / "invalid.json"
        invalid.write_text('{"items": [broken}', encoding="utf-8")
        self.assertEqual(self.extract_unchanged(invalid)["text"], '{"items": [broken}')

    def test_source_path_cannot_be_output_directory(self) -> None:
        source = self.root / "input.txt"
        source.write_text("Original text", encoding="utf-8")
        before = digest(source)
        with self.assertRaises(ValueError):
            extract_source(source, source)
        self.assertEqual(before, digest(source))
        another_file = self.root / "output.txt"
        another_file.write_text("Do not overwrite", encoding="utf-8")
        with self.assertRaises(ValueError):
            extract_source(source, another_file)
        self.assertEqual(another_file.read_text(encoding="utf-8"), "Do not overwrite")

    def test_real_docx_conversion(self) -> None:
        from docx import Document

        document = Document()
        document.add_paragraph("1. 【填空】小数与单词？")
        document.add_paragraph("正确答案：1.5 and two words")
        document.add_paragraph("解析：第一行解释")
        document.add_paragraph("第二行解释")
        source = self.root / "questions.docx"
        document.save(source)
        result = self.extract_unchanged(source)
        self.assertIn("1.5 and two words", result["text"])
        self.assertIn("第二行解释", result["text"])
        self.assertFalse(result["needsVisualReview"])
        self.assertTrue(Path(result["mediaDir"]).is_relative_to(self.work.resolve()))
        parsed = parse_question_text(result["text"])
        self.assertEqual(len(parsed["items"]), 1, result["text"])
        self.assertEqual(parsed["items"][0]["correctAnswer"], ["1.5 and two words"])
        self.assertIn("第二行解释", parsed["items"][0]["originalExplanation"])

    def test_docx_image_is_extracted_but_not_claimed_read(self) -> None:
        from docx import Document
        from PIL import Image

        picture = self.root / "diagram.png"
        Image.new("RGB", (40, 30), "white").save(picture)
        document = Document()
        document.add_paragraph("1.【简答】解释图片。")
        document.add_picture(str(picture))
        source = self.root / "with_image.docx"
        document.save(source)
        result = self.extract_unchanged(source)
        self.assertTrue(result["needsVisualReview"])
        self.assertTrue(list(Path(result["mediaDir"]).rglob("*.png")))
        self.assertTrue(parse_question_text(result["text"])["items"][0]["images"])

    def test_docx_choice_paragraphs_keep_options(self) -> None:
        from docx import Document

        document = Document()
        for paragraph in ["1. 【单选】选一个。", "A. 甲", "B. 乙", "答案：B"]:
            document.add_paragraph(paragraph)
        source = self.root / "choices.docx"
        document.save(source)
        result = self.extract_unchanged(source)
        parsed = parse_question_text(result["text"])
        self.assertEqual(len(parsed["items"]), 1, result["text"])
        self.assertEqual(parsed["items"][0]["options"], [{"label": "A", "text": "甲"}, {"label": "B", "text": "乙"}])
        self.assertEqual(parsed["items"][0]["correctAnswer"], ["B"])

    def test_docx_personal_image_removed_only_from_derived_media(self) -> None:
        from docx import Document
        from PIL import Image

        picture = self.root / "personal.png"
        Image.new("RGB", (30, 30), "red").save(picture)
        picture_before = digest(picture)
        document = Document()
        document.add_paragraph("1. 【简答】原题没有图片。")
        document.add_paragraph("正确答案：original")
        document.add_paragraph("你的答案：PRIVATE")
        document.add_picture(str(picture))
        document.add_paragraph("PRIVATE_MORE")
        document.add_paragraph("解析：原解释")
        source = self.root / "personal_image.docx"
        document.save(source)
        result = self.extract_unchanged(source)
        self.assertEqual(digest(picture), picture_before)
        self.assertNotIn("PRIVATE", result["text"])
        self.assertFalse(result["needsVisualReview"])
        self.assertEqual(list(Path(result["mediaDir"]).rglob("*.png")), [])
        item = parse_question_text(result["text"])["items"][0]
        self.assertEqual(item["images"], [])
        self.assertEqual(item["correctAnswer"], ["original"])

    def test_docx_shared_question_and_personal_image_is_retained(self) -> None:
        from docx import Document
        from PIL import Image

        picture = self.root / "shared.png"
        Image.new("RGB", (30, 30), "blue").save(picture)
        document = Document()
        document.add_paragraph("1. 【简答】见题图。")
        document.add_picture(str(picture))
        document.add_paragraph("正确答案：original")
        document.add_paragraph("我的答案：PRIVATE")
        document.add_picture(str(picture))
        source = self.root / "shared_image.docx"
        document.save(source)
        result = self.extract_unchanged(source)
        self.assertNotIn("PRIVATE", result["text"])
        self.assertTrue(result["needsVisualReview"])
        self.assertEqual(len(list(Path(result["mediaDir"]).rglob("*.png"))), 1)
        self.assertEqual(len(parse_question_text(result["text"])["items"][0]["images"]), 1)

    def test_real_pdf_keeps_page_boundaries(self) -> None:
        from reportlab.pdfgen import canvas

        source = self.root / "questions.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawString(20, 750, "1. [fill] What decimal?")
        pdf.drawString(20, 730, "Answer: 1.5")
        pdf.showPage()
        pdf.drawString(20, 750, "2. [short] How many words?")
        pdf.drawString(20, 730, "Answer: two words")
        pdf.save()
        result = self.extract_unchanged(source)
        self.assertIn("<!-- PDF page 1 -->", result["text"])
        self.assertIn("<!-- PDF page 2 -->", result["text"])
        self.assertFalse(result["needsVisualReview"])
        self.assertEqual(result["needsVisualPages"], [])
        questions = parse_question_text(result["text"])["items"]
        self.assertEqual(len(questions), 2)
        self.assertEqual(questions[0]["correctAnswer"], ["1.5"])
        self.assertEqual(questions[1]["correctAnswer"], ["two words"])

    def test_scanned_pdf_and_image_require_visual_transcription(self) -> None:
        from PIL import Image
        from reportlab.pdfgen import canvas

        picture = self.root / "scan.png"
        Image.new("RGB", (100, 100), "white").save(picture)
        image_result = self.extract_unchanged(picture)
        self.assertEqual(image_result["text"], "")
        self.assertTrue(image_result["needsVisualReview"])
        source = self.root / "scanned.pdf"
        pdf = canvas.Canvas(str(source))
        pdf.drawImage(str(picture), 20, 20)
        pdf.save()
        result = self.extract_unchanged(source)
        self.assertTrue(result["needsVisualReview"])
        self.assertEqual(result["needsVisualPages"], [1])
        self.assertTrue(any("no extractable text" in note for note in result["notes"]))
        self.assertTrue(any("embedded images" in note for note in result["notes"]))
        self.assertEqual(parse_question_text(result["text"])["items"], [])

    def test_legacy_doc_rejected_without_modification(self) -> None:
        source = self.root / "legacy.doc"
        source.write_bytes(b"not-a-real-doc")
        before = digest(source)
        with self.assertRaisesRegex(ValueError, "DOCX"):
            extract_source(source, self.work)
        self.assertEqual(before, digest(source))


class ParserTests(unittest.TestCase):
    def test_indeterminate_and_qa_sections_with_lowercase_options(self) -> None:
        text = """# 任选题
1. 哪些成立？
a. 第一项
b、第二项
c) 第三项
答案：a/c
# 问答题
2. 解释结论。
答案：完整的多行答案
第二行
"""
        items = parse_question_text(text)["items"]
        self.assertEqual([item["type"] for item in items], ["不定项", "简答"])
        self.assertEqual([option["label"] for option in items[0]["options"]], ["A", "B", "C"])
        self.assertEqual(items[0]["correctAnswer"], ["A", "C"])
        self.assertEqual(items[1]["correctAnswer"], ["完整的多行答案\n第二行"])
        for alias in ("不定项", "不定项选择题", "任选", "问答"):
            with self.subTest(alias=alias):
                item = parse_question_text(f"1.【{alias}】题干？\n答案：文字")["items"][0]
                self.assertEqual(item["type"], "简答" if alias == "问答" else "不定项")

    def test_pending_answer_and_reference_remain_unconfirmed_after_reintake(self) -> None:
        for status in ("待核实", "未提供，待核实", "待核实（原题参考答案：完整答案）",
                       "待核实(参考答案:完整答案)"):
            for kind in ("填空", "问答"):
                with self.subTest(status=status, kind=kind):
                    parsed = parse_question_text(f"1.【{kind}】题干？\n答案：{status}\n解析：保留解析。")
                    item = parsed["items"][0]
                    self.assertEqual(item["correctAnswer"], [])
                    self.assertEqual(item["originalAnswer"], status)
                    self.assertIn(status, item["rawText"])
                    self.assertTrue(any(issue["code"] == "requires_content_review" for issue in item["issues"]))
        literal = parse_question_text("1.【问答】解释状态。\n答案：待核实是未完成核验的状态。")["items"][0]
        self.assertEqual(literal["correctAnswer"], ["待核实是未完成核验的状态。"])

    def test_numbered_qa_steps_larger_than_question_number_are_retained(self) -> None:
        text = "1.【问答】描述过程。\n答案：\n1. 先确定条件\n2. 再计算结果\n3. 最后检查\n4.【填空】结果？\n答案：1.5"
        parsed = parse_question_text(text)
        self.assertEqual(len(parsed["items"]), 2)
        self.assertEqual(parsed["items"][0]["correctAnswer"], ["1. 先确定条件\n2. 再计算结果\n3. 最后检查"])
        self.assertTrue(any(issue["code"] == "numbered_field_continuation" for issue in parsed["items"][0]["issues"]))
        self.assertEqual(parsed["items"][1]["correctAnswer"], ["1.5"])
        explicit = parse_question_text("1.【问答】解释。\n答案：回答\n2. 下一题？\n答案：下一题答案")
        self.assertEqual(len(explicit["items"]), 2)
        self.assertTrue(any(warning["code"] == "ambiguous_question_boundary" for warning in explicit["warnings"]))

    def test_existing_case_and_field_specific_media_are_preserved(self) -> None:
        text = """1.【问答】![题图](media/shared.png)
案例：保留已有案例条件。
后续条件。
答案：![答案图](media/answer.png)
解析：![解析图](media/shared.png)
2.【任选】看图选择？
a. ![甲图](media/a.png)
b. ![乙图](media/b.png)
答案：ab
"""
        first, second = parse_question_text(text)["items"]
        self.assertEqual(first["case"], "保留已有案例条件。\n后续条件。")
        self.assertEqual([image["field"] for image in first["images"]], ["title", "answer", "explanation"])
        self.assertEqual([image["path"] for image in first["images"]],
                         ["media/shared.png", "media/answer.png", "media/shared.png"])
        self.assertEqual([(image["field"], image["optionIndex"]) for image in second["images"]],
                         [("option", 0), ("option", 1)])
        self.assertEqual(second["correctAnswer"], ["A", "B"])

    def test_supported_types_option_syntax_and_multiline_fields(self) -> None:
        text = """课程说明，保留追溯。
1.【单选题】请选择。
A. 选项一
B、选项二
正确答案：B
用户答案：A
解析：第一行
第二行
二、【多选】哪些成立？
A) 第一项
B. 第二项
C：第三项
答案：AC
我的答案：A、C
3.【判断】这个说法正确吗？
答案：正确
（4）【填空】写出值和词组。
答案：1.5
two words
Q5: [short answer] Explain.
Answer: line one
line two
Explanation: detail one
detail two
"""
        result = parse_question_text(text)
        self.assertEqual([item["type"] for item in result["items"]], ["单选", "多选", "判断", "填空", "简答"])
        self.assertEqual(len(result["unparsed"]), 1)
        first, second, _, fourth, fifth = result["items"]
        self.assertEqual(first["options"], [{"label": "A", "text": "选项一"}, {"label": "B", "text": "选项二"}])
        self.assertNotIn("selectedAnswer", first)
        self.assertEqual(first["originalExplanation"], "第一行\n第二行")
        self.assertEqual(second["correctAnswer"], ["A", "C"])
        self.assertNotIn("selectedAnswer", second)
        self.assertEqual(fourth["correctAnswer"], ["1.5\ntwo words"])
        self.assertEqual(fifth["correctAnswer"], ["line one\nline two"])
        self.assertEqual(fifth["originalExplanation"], "detail one\ndetail two")

    def test_missing_answers_and_unknown_types_are_not_guessed(self) -> None:
        result = parse_question_text("1.【计算题】求值。\nA. 1\nB. 2\n2. 无题型。\n答案：two words")
        self.assertEqual(len(result["items"]), 2)
        first, second = result["items"]
        self.assertEqual(first["rawType"], "计算题")
        self.assertEqual(first["type"], "未知")
        self.assertEqual(first["correctAnswer"], [])
        self.assertEqual(second["type"], "未知")
        self.assertEqual(second["correctAnswer"], ["two words"])
        self.assertTrue(any(warning["code"] == "missing_answer" for warning in result["warnings"]))
        self.assertEqual(sum(warning["code"] == "unknown_type" for warning in result["warnings"]), 2)

    def test_decimal_lines_are_not_question_headers(self) -> None:
        result = parse_question_text("1.【填空】值？\n答案：\n1.5\ntwo words\n解析：结果为1.5。")
        self.assertEqual(len(result["items"]), 1)
        self.assertEqual(result["items"][0]["correctAnswer"], ["1.5\ntwo words"])

    def test_explicit_section_type_and_standalone_answer_marker(self) -> None:
        result = parse_question_text("【填空题】\n1.数值？\n答案\n1.5\n2.词组？\n正确答案：two words")
        self.assertEqual([item["type"] for item in result["items"]], ["填空", "填空"])
        self.assertEqual(result["items"][0]["correctAnswer"], ["1.5"])
        self.assertEqual(result["items"][1]["correctAnswer"], ["two words"])

    def test_numbered_answer_steps_kept_and_ambiguous_boundaries_warn(self) -> None:
        text = "7.【简答】过程？\n答案：第一步\n1. 先判断\n2. 再计算\n8. 另一个问题？\n答案：回答"
        result = parse_question_text(text)
        self.assertEqual(len(result["items"]), 2)
        self.assertIn("1. 先判断\n2. 再计算", result["items"][0]["correctAnswer"][0])
        self.assertTrue(any(w["code"] == "ambiguous_question_boundary" for w in result["warnings"]))
        self.assertTrue(all(any(issue["code"] == "ambiguous_question_boundary" for issue in item["issues"]) for item in result["items"]))
        self.assertIn("2. 再计算", result["items"][0]["rawText"])

    def test_pdf_page_marker_is_source_locator(self) -> None:
        result = parse_question_text("<!-- PDF page 3 -->\n1.【判断】说法？\n答案：正确")
        self.assertEqual(result["items"][0]["sourcePage"], 3)
        self.assertEqual(result["items"][0]["sourcePages"], [3])

    def test_question_before_scan_page_is_conservatively_associated(self) -> None:
        text = "<!-- PDF page 1 -->\n1.【简答】见下图。\n答案：回答\n<!-- PDF page 2 -->\n\n<!-- PDF page 3 -->\n2.【填空】文字？\n答案：word"
        result = parse_question_text(text)
        self.assertEqual(result["items"][0]["sourcePages"], [1, 2])
        self.assertEqual(result["items"][1]["sourcePages"], [3])

    def test_id_is_stable_independent_of_number_and_answer(self) -> None:
        first = parse_question_text("1.【单选】同题？\nA. 甲\nB. 乙\n答案：A")["items"][0]
        second = parse_question_text("9.【单选】同题？\nA. 甲\nB. 乙\n答案：B")["items"][0]
        self.assertEqual(first["id"], second["id"])

    def test_unrecognized_text_is_reported_and_duplicate_labels_warn(self) -> None:
        result = parse_question_text("无法判定格式的文字\n完整保留。")
        self.assertEqual(result["items"], [])
        self.assertIn("完整保留", result["unparsed"][0]["text"])
        duplicate = parse_question_text("1.【单选】重复？\nA. 甲\nA. 乙\n答案：A")
        self.assertEqual(len(duplicate["items"][0]["options"]), 2)
        self.assertTrue(any(w["code"] == "duplicate_option_label" for w in duplicate["warnings"]))
        self.assertTrue(any(w["code"] == "duplicate_option_label" for w in duplicate["items"][0]["issues"]))

    def test_images_keep_original_references(self) -> None:
        item = parse_question_text('1.【简答】![图](media/figure.png)\n答案：文字\n解析：<img src="media/second.png">')["items"][0]
        self.assertEqual(item["images"], [{"path": "media/figure.png", "alt": "图", "required": True, "field": "title"}, {"path": "media/second.png", "alt": "", "required": True, "field": "explanation"}])

    def test_personal_multiline_decimal_and_numbered_steps_are_removed(self) -> None:
        text = """7.【简答】描述过程。
正确答案：原始完整答案
用户答案：PRIVATE_INLINE
PRIVATE_CONTINUATION
1. PRIVATE_STEP_ONE
2. PRIVATE_STEP_TWO
1.5
two personal words
解析：原始解析
第二行原解析
8.【判断】下一题？
答案：正确
"""
        cleaned = strip_personal_answer_text(text)
        self.assertEqual(cleaned.count("\n"), text.count("\n"))
        self.assertNotIn("PRIVATE", cleaned)
        self.assertNotIn("two personal words", cleaned)
        self.assertEqual(cleaned, strip_personal_answer_text(cleaned))
        result = parse_question_text(text)
        self.assertEqual(len(result["items"]), 2)
        self.assertEqual(result["items"][0]["correctAnswer"], ["原始完整答案"])
        self.assertEqual(result["items"][0]["originalExplanation"], "原始解析\n第二行原解析")
        self.assertEqual(result["items"][1]["sourceLine"], 11)
        serialized = json.dumps(result, ensure_ascii=False)
        self.assertNotIn("PRIVATE", serialized)
        self.assertNotIn("selectedAnswer", serialized)

    def test_personal_labels_and_bracket_fields_are_removed(self) -> None:
        labels = ["我的答案", "个人作答", "用户答案", "你的答案", "学生答案", "个人答案",
                  "Your  Answer", "My Answer", "selected answer"]
        for label in labels:
            with self.subTest(label=label):
                text = f"1.【简答】问题？\n{label}：PRIVATE\nPRIVATE_MORE\n正确答案：original"
                parsed = parse_question_text(text)
                self.assertEqual(parsed["items"][0]["correctAnswer"], ["original"])
                self.assertNotIn("PRIVATE", json.dumps(parsed, ensure_ascii=False))
        for marker in ("【我的答案】PRIVATE", "[your answer] PRIVATE", "**用户答案**：PRIVATE", "# 个人作答：PRIVATE"):
            with self.subTest(marker=marker):
                text = "1.【简答】问题？\n" + marker + "\nPRIVATE_MORE\n解析：原解释"
                self.assertNotIn("PRIVATE", json.dumps(parse_question_text(text), ensure_ascii=False))

    def test_personal_answer_never_fills_missing_correct_answer(self) -> None:
        result = parse_question_text("1.【填空】原件没有正确答案？\n我的答案：PRIVATE\nPRIVATE_MORE")
        item = result["items"][0]
        self.assertEqual(item["correctAnswer"], [])
        self.assertEqual(item["originalAnswer"], "")
        self.assertNotIn("selectedAnswer", item)
        self.assertNotIn("PRIVATE", item["rawText"])
        self.assertTrue(any(w["code"] == "missing_answer" for w in result["warnings"]))

    def test_personal_answer_image_does_not_become_question_image(self) -> None:
        text = """1.【简答】![题图](media/question.png)
正确答案：原答案
你的答案：![PRIVATE](media/personal.png)
<img src="media/personal-more.png">
解析：![解析图](media/explanation.png)
"""
        item = parse_question_text(text)["items"][0]
        self.assertEqual([image["path"] for image in item["images"]], ["media/question.png", "media/explanation.png"])
        self.assertNotIn("personal.png", item["rawText"])
        self.assertNotIn("personal-more.png", item["rawText"])

    def test_personal_unparsed_text_is_not_retained(self) -> None:
        text = "课程说明\n我的答案：PRIVATE\nPRIVATE_MORE\n1.【简答】问题？\n答案：原答案"
        result = parse_question_text(text)
        self.assertEqual(result["unparsed"][0]["text"], "课程说明")
        self.assertNotIn("PRIVATE", json.dumps(result, ensure_ascii=False))

    def test_plain_numbered_next_question_with_options_stops_personal_field(self) -> None:
        text = """【单选题】
1. 第一题
A. 甲
B. 乙
我的答案：PRIVATE
PRIVATE_MORE
2. 第二题
A. 丙
B. 丁
正确答案：B
"""
        result = parse_question_text(text)
        self.assertEqual(len(result["items"]), 2)
        self.assertEqual(result["items"][0]["correctAnswer"], [])
        self.assertEqual(result["items"][1]["title"], "第二题")
        self.assertEqual(result["items"][1]["correctAnswer"], ["B"])
        self.assertNotIn("PRIVATE", json.dumps(result, ensure_ascii=False))

    def test_unclear_unlabeled_text_and_literal_mentions_are_not_deleted(self) -> None:
        text = "普通讨论：我的答案与原书不同。\n未标记为个人答案的正文必须保留。\n1.【简答】解释“用户答案”这个字段。\n答案：这是显示名称。"
        self.assertEqual(strip_personal_answer_text(text), text)
        result = parse_question_text(text)
        self.assertIn("未标记为个人答案", result["unparsed"][0]["text"])
        self.assertIn("用户答案", result["items"][0]["title"])

    def test_cleaning_retains_crlf_and_pdf_provenance(self) -> None:
        text = "<!-- PDF page 1 -->\r\n1.【简答】问题？\r\n用户答案：PRIVATE\r\n<!-- PDF page 2 -->\r\nPRIVATE_MORE\r\n正确答案：original\r\n"
        cleaned = strip_personal_answer_text(text)
        self.assertEqual(cleaned.count("\r\n"), text.count("\r\n"))
        self.assertIn("<!-- PDF page 2 -->", cleaned)
        item = parse_question_text(text)["items"][0]
        self.assertEqual(item["sourcePages"], [1, 2])
        self.assertEqual(item["sourceLine"], 2)
        self.assertEqual(item["correctAnswer"], ["original"])


if __name__ == "__main__":
    unittest.main()
