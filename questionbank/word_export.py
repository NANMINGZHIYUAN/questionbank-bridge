"""Fojiao Word export from reviewed records, without generating answers."""
from __future__ import annotations

import json
import hashlib
from collections import Counter
from pathlib import Path
import re
import string
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from .core import strip_personal_answers, validate_bank
from .exporters import _answers, _comparison, _markdown, _reference_document, _text, _type


HEADINGS = {"单选": "单选题", "多选": "多选题", "不定项": "不定项选择题",
            "判断": "判断题", "填空": "填空题", "简答": "问答题"}
CHOICES = {"单选", "多选", "不定项"}
CONTENT_CODES = {
    "missing_correct_answer", "unmapped_answer", "ambiguous_answer",
    "conflicting_correct_answers", "unknown_type", "missing_required_image",
    "requires_visual_review", "requires_content_review", "ambiguous_question_boundary",
    "numbered_field_continuation", "duplicate_option_label", "empty_title",
    "duplicate_options", "invalid_answer_cardinality", "invalid_image_placement", "invalid_required_image",
}
TRUTHS = {"正确", "对", "√", "是", "true", "t", "1", "y"}
FALSEHOODS = {"错误", "错", "×", "否", "false", "f", "0", "n"}
PRIVATE_PATH = re.compile(r"(?i)(?:file://[^\s<>]+|(?<![a-z0-9])(?:[a-z]:[\\/]|\\\\)[^\s<>]+)")


class FormatManual(ValueError):
    """Verified content that cannot be represented by the Word import format."""


class MissingRenderMedia(RuntimeError):
    def __init__(self, filenames):
        self.filenames = filenames
        super().__init__('生成的 Word 缺少应有图片：' + '、'.join(filenames))


def _clean(value):
    # Use existing citation filtering, then remove private bare paths and learner
    # response lines as well. This never changes registered source files.
    text = str(value or '')
    if re.search(r'<table\b', text, re.I):
        from .manual_text import _privacy_fallback
        text = _privacy_fallback(text, [])
    text = _markdown(text)
    text = strip_personal_answers({'title': text})['title']
    text = PRIVATE_PATH.sub("[本地路径已省略]", text)
    # Blank lines such as ____ are content, rather than Markdown emphasis.
    # Math, code and hyperlink destinations retain their own underscore syntax.
    pieces = re.split(r"(```.*?```|~~~.*?~~~|`[^`\n]+`|\$\$.*?\$\$|(?<!\$)\$[^$\n]+\$(?!\$)|!?\[[^\]\n]*\]\([^\n]*?\))", text, flags=re.S)
    for index in range(0, len(pieces), 2):
        pieces[index] = re.sub(r"(?<![\\_])_{2,}", lambda match: r"\_" * len(match[0]), pieces[index])
    return ''.join(pieces)


def _choice_answer(question):
    options = question.get("options", [])
    if not options:
        raise ValueError("选择题选项为空，需核对原题。")
    texts = [_comparison(o.get("text", "")) for o in options]
    labels = [_comparison(o.get("label", "")).upper() for o in options]
    if len(set(texts)) != len(texts) or len(set(labels)) != len(labels):
        raise ValueError("选项内容或标签重复，答案无法唯一对应。")
    selected = []
    for answer in _answers(question):
        value = _comparison(answer)
        matches = [i for i, text in enumerate(texts) if text == value]
        # Canonical correctAnswer contains option text. Legacy banks may still
        # contain labels; accept that only when no option text matches.
        if not matches:
            matches = [i for i, label in enumerate(labels) if label == value.upper()]
        if len(matches) != 1:
            raise ValueError("正确答案无法唯一对应当前选项。")
        selected.append(matches[0])
    if not selected or len(set(selected)) != len(selected):
        raise ValueError("选择题缺答案或答案重复。")
    if _type(question) == "单选" and len(selected) != 1:
        raise ValueError("单选题应有且仅有一个正确选项。")
    if len(options) > 26:
        raise FormatManual("option_count_exceeds_word_labels")
    return "".join(string.ascii_uppercase[i] for i in sorted(selected))


def _judgment_answer(question):
    answers = _answers(question)
    if len(answers) != 1:
        raise ValueError("判断题应有且仅有一个正确答案。")
    value = _comparison(answers[0]).lower()
    if value not in TRUTHS | FALSEHOODS:
        matches = [o for o in question.get("options", [])
                   if _comparison(o.get("label", "")).lower() == value]
        if len(matches) == 1:
            value = _comparison(matches[0].get("text", "")).lower()
    if value in TRUTHS:
        return "正确"
    if value in FALSEHOODS:
        return "错误"
    raise ValueError("判断答案不能唯一转换为正确或错误。")


def _answer(question):
    kind = _type(question)
    if kind in CHOICES:
        return _choice_answer(question)
    if kind == "判断":
        return _judgment_answer(question)
    # Word accepts a separate answer field; do not invent JSON-style {answer}
    # blanks or split answer text on punctuation.
    return ("；" if kind == "填空" else "\n\n").join(_answers(question))


def _brief(question, report):
    if question.get("briefExplanation", "").strip():
        return _clean(question["briefExplanation"])
    available = [question.get("originalExplanation", ""), question.get("explanation", "")]
    blocks = []
    for value in available:
        text = _clean(value).strip()
        if text and text not in blocks:
            blocks.append(text)
    report["wordExport"]["needsBriefExplanation"].append(question["id"])
    warning = f"题目 {question['id']} 未写简析；已有可靠解析完整保留，无解析时留空，待人工简要整理。"
    report["warnings"].append(warning)
    report["wordExport"]["warnings"].append(warning)
    return "\n\n".join(blocks)


def _image_groups(question, media, report):
    groups = {"title": [], "answer": [], "explanation": []}
    groups["option"] = {}
    qid = question["id"]
    available = list(media.get(qid, []))
    used = []
    for item in question.get("images", []):
        path = item.get("path")
        filename = Path(path).name if path else None
        if filename not in available:
            continue
        field = item.get("field", "title")
        if "field" not in item:
            warning = f"题目 {qid} 的旧图片无位置标注，暂置于题干后；请按原件核对位置。"
            if warning not in report["wordExport"]["warnings"]:
                report["wordExport"]["warnings"].append(warning)
        if field == "option":
            index = item.get("optionIndex")
            if not isinstance(index, int) or isinstance(index, bool) or not 0 <= index < len(question.get("options", [])):
                raise ValueError("选项图片的位置索引无效，需回看原题。")
            groups["option"].setdefault(index, []).append(filename)
        elif field in groups:
            groups[field].append(filename)
        else:
            raise ValueError("题图位置字段无效，需回看原题。")
        used.append(filename)
    # Support callers passing media without mutated image paths, as well as old
    # records. Omitted optional images do not shift later option associations.
    for filename in available:
        if filename not in used:
            groups["title"].append(filename)
            warning = f"题目 {qid} 的旧图片无位置标注，暂置于题干后；请按原件核对位置。"
            if warning not in report['wordExport']['warnings']:
                report['wordExport']['warnings'].append(warning)
    return groups


def _images(files):
    return [f"![题图](assets/{filename})" for filename in files]


def _check_media(question, media, out, cache):
    """Decode local raster assets; preserve unsupported readable media for manual handling."""
    from PIL import Image
    problems = []
    for filename in set(media.get(question['id'], [])):
        if filename not in cache:
            path = out / 'assets' / filename
            try:
                if path.suffix.lower() == '.svg':
                    root = ET.parse(path).getroot()
                    if root.tag.rsplit('}', 1)[-1] != 'svg':
                        raise ValueError('not SVG')
                    cache[filename] = 'unsupported_vector_media'
                else:
                    with Image.open(path) as image:
                        image.verify()
                    with Image.open(path) as image:
                        image.load()
                    cache[filename] = None
            except (OSError, ValueError, SyntaxError, ET.ParseError, Image.DecompressionBombError):
                cache[filename] = 'unreadable_media'
        if cache[filename]:
            required = any(Path(item.get('path') or '').name == filename and item.get('required', False)
                           for item in question.get('images', []))
            problems.append({'file': 'assets/' + filename, 'reason': cache[filename], 'required': bool(required)})
    return problems


def _option_label(option, index, preserve=False):
    if preserve:
        return str(option.get('label') or index + 1).upper()
    return string.ascii_uppercase[index] if index < 26 else str(option.get("label") or index + 1)


def _manual_answer(question):
    if _type(question) not in CHOICES or len(question.get('options', [])) <= 26:
        return _answer(question)
    options = question['options']
    texts = [_comparison(option.get('text', '')) for option in options]
    labels = [_comparison(option.get('label', '')).upper() for option in options]
    answers = []
    for answer in _answers(question):
        value = _comparison(answer)
        matches = [index for index, text in enumerate(texts) if text == value]
        if not matches:
            matches = [index for index, label in enumerate(labels) if label == value.upper()]
        if len(matches) != 1:
            raise ValueError('手动答案无法唯一对应当前选项。')
        answers.append(options[matches[0]]['text'])
    return '\n\n'.join(answers)


def _write_manual_materials(material, out, pandoc, report, media):
    """Keep the source Markdown and produce separate plain-text copy fields."""
    from .manual_text import render_plain
    markdown = ["# 佛脚刷题完整手动材料", "", "此文件保留原公式、表格和附件，不直接粘贴 Markdown 到平台文字框。", ""]
    plain = ["佛脚刷题手动录入纯文本", "", "按栏复制到对应输入框；栏目标题、题目编号及附件说明不需要粘贴。",
             "空案例栏保持为空。常见公式和表格转为可读文字；有格式告警的内容须对照原稿核对。",
             "图片附件不会随文字自动上传。关键图在平台缺失时，该题不能视为完整导入。", ""]
    for number, question in enumerate(material, 1):
        qid = question['id']
        markdown += _question_lines(question, number, report, media, manual=True)
        fields = [('案例', question.get('case', '')), ('题干', question.get('title', ''))]
        preserve = len(question.get('options', [])) > 26
        fields += [('选项' + _option_label(option, index, preserve), option.get('text', ''))
                   for index, option in enumerate(question.get('options', []))]
        answer = _manual_answer(question)
        fields += [('答案', answer), ('解析', question.get('_wordBrief', ''))]
        plain += [f"题目 {number}（{HEADINGS.get(_type(question), _type(question))}）", ""]
        for label, value in fields:
            text, warnings = render_plain(_clean(value), pandoc)
            plain += [f"【{label}】", text, ""]
            if warnings:
                plain += [f"【{label}格式需核对，不粘贴到文字框】", *warnings, ""]
            for warning in warnings:
                entry = {'questionId':qid, 'field':label, 'message':warning}
                if entry not in report['wordExport']['manualTextWarnings']:
                    report['wordExport']['manualTextWarnings'].append(entry)
        attachments = []
        for image in question.get('images', []):
            filename = Path(image.get('path', '')).name
            if filename not in media.get(qid, []):
                continue
            field = image.get('field', 'title')
            location = {'title':'题干', 'option':'选项', 'answer':'答案', 'explanation':'解析'}.get(field, field)
            if field == 'option':
                location += _option_label(question['options'][image['optionIndex']], image['optionIndex'], preserve)
            attachments.append(f"{location}：assets/{filename}（{'关键图' if image.get('required') else '非关键图'}）")
            entry = {'questionId':qid, 'field':field, 'file':'assets/'+filename,
                     'required':bool(image.get('required')), 'platformDisplayVerified':False}
            if field == 'option':
                entry['optionIndex'] = image['optionIndex']
            if entry not in report['wordExport']['mediaAttachments']:
                report['wordExport']['mediaAttachments'].append(entry)
        declared = {Path(item.get('path') or '').name for item in question.get('images', [])}
        for filename in media.get(qid, []):
            if filename in declared:
                continue
            attachments.append(f"题干后：assets/{filename}（旧图位置未标注，需按原件核对）")
            entry = {'questionId':qid, 'field':'title', 'file':'assets/'+filename,
                     'required':False, 'platformDisplayVerified':False, 'placementNeedsReview':True}
            if entry not in report['wordExport']['mediaAttachments']:
                report['wordExport']['mediaAttachments'].append(entry)
        if attachments:
            plain += ['【附件位置说明，不粘贴到文字框】', *attachments, '']
        plain += ['────────', '']
    source = out / '问答题手动录入.md'
    source.write_text('\n'.join(markdown).rstrip()+'\n', encoding='utf-8')
    target = out / '问答题手动录入.txt'
    target.write_text('\n'.join(plain).rstrip()+'\n', encoding='utf-8')
    report['files']['manualQuestions'] = str(target)
    report['files']['manualQuestionsMarkdown'] = str(source)


def _question_lines(question, number, report, media, manual=False):
    kind = _type(question)
    groups = _image_groups(question, media, report)
    title = _clean(question.get("title", ""))
    case = _clean(question.get("case", ""))
    if manual:
        lines = [f"## {number} {HEADINGS.get(kind, kind)}", "", "### 案例", "", case, "",
                 "### 题干", "", title, ""]
    else:
        stem = ("案例：" + case + "\n\n" if case else "") + title
        lines = [f"{number}、" + stem, ""]
    lines += _images(groups["title"]) + ([""] if groups["title"] else [])
    if kind in CHOICES:
        for i, option in enumerate(question.get("options", [])):
            lines += [_clean(_option_label(option, i, manual and len(question['options']) > 26)) + r"\. " + _clean(option.get("text", "")), ""]
            lines += _images(groups["option"].get(i, [])) + ([""] if groups["option"].get(i) else [])
    elif groups["option"]:
        # Nonstandard judgment options must keep their images and identity.
        for i, option in enumerate(question.get("options", [])):
            lines += [_clean(_option_label(option, i)) + r"\. " + _clean(option.get("text", "")), ""]
            lines += _images(groups["option"].get(i, [])) + [""]
    if manual:
        answer = _manual_answer(question)
        lines += ["### 答案", "", _clean(answer), ""]
    else:
        lines += ["答案：" + _clean(_answer(question)), ""]
    lines += _images(groups["answer"]) + ([""] if groups["answer"] else [])
    brief = question.get("_wordBrief", "")
    lines += (["### 解析", "", brief, ""] if manual else ["解析：" + brief, ""])
    lines += _images(groups["explanation"]) + ([""] if groups["explanation"] else [])
    return lines


def _render_docx(questions, target, out, pandoc, reference, report):
    from docx import Document
    lines, previous = [], None
    for number, question in enumerate(questions, 1):
        kind = _type(question)
        if kind != previous:
            lines += ["# " + HEADINGS[kind], ""]
            previous = kind
        lines += _question_lines(question, number, report, report["_wordMedia"])
    source = target.with_suffix(".md")
    source.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    command = [str(pandoc), str(source), "--from", "markdown+tex_math_dollars+tex_math_single_backslash-raw_html-implicit_figures",
               "--to", "docx", "--standalone", "--resource-path", str(out),
               "--reference-doc", str(reference), "--output", str(target)]
    result = subprocess.run(command, cwd=out, capture_output=True, text=True, encoding="utf-8", timeout=180)
    if result.returncode:
        raise RuntimeError("Pandoc 生成题库 Word 失败：" + result.stderr.strip())
    if result.stderr.strip():
        report["warnings"].append("Pandoc 题库: " + result.stderr.strip())
    document = Document(target)
    # Do not rediscover question boundaries from numbered answer steps. The
    # split boundaries are the structured records passed into this function.
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.keep_together = True
        if paragraph.style.name.startswith("Heading") or re.match(r"^[A-Z]\. ", paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
    document.save(target)
    # Count actual drawing references, including repeated uses of the same
    # image in different fields. An unused ZIP media member is insufficient.
    expected = Counter()
    filenames = {}
    for question in questions:
        groups = _image_groups(question, report['_wordMedia'], report)
        locations = groups['title'] + groups['answer'] + groups['explanation']
        locations += [filename for images in groups['option'].values() for filename in images]
        for filename in locations:
            digest = hashlib.sha256((out / 'assets' / filename).read_bytes()).hexdigest()
            expected[digest] += 1
            filenames.setdefault(digest, set()).add(filename)
    with zipfile.ZipFile(target) as package:
        if package.testzip():
            raise RuntimeError('生成的 Word ZIP 校验失败。')
        if expected:
            root = ET.fromstring(package.read('word/document.xml'))
            relationships = ET.fromstring(package.read('word/_rels/document.xml.rels'))
            targets = {item.get('Id'): item.get('Target') for item in relationships}
            actual = Counter()
            namespaces = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
            for drawing in root.findall('.//a:blip', namespaces):
                relative = targets.get(drawing.get('{' + namespaces['r'] + '}embed'))
                if relative:
                    member = 'word/' + relative.lstrip('/')
                    actual[hashlib.sha256(package.read(member)).hexdigest()] += 1
            missing = [filename for digest, count in expected.items() if actual[digest] < count
                       for filename in filenames[digest]]
            if missing:
                raise MissingRenderMedia(sorted(set(missing)))


def _compact_docx(path):
    """Recompress ZIP members without altering content, images or equations."""
    compacted = path.with_suffix(".compact.docx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(compacted, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for member in source.infolist():
            member.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(member, source.read(member.filename), compresslevel=9)
    if compacted.stat().st_size < path.stat().st_size:
        compacted.replace(path)
    else:
        compacted.unlink()


def export_word(bank, out, pandoc, report, media, required_missing, word_qa_mode="manual", max_bytes=500000):
    if word_qa_mode not in {"include", "manual"}:
        raise ValueError("word_qa_mode 必须为 include 或 manual。")
    if not isinstance(max_bytes, int) or isinstance(max_bytes, bool) or max_bytes <= 0:
        raise ValueError("Word 文件大小上限必须是正整数字节数。")
    out = Path(out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    working = strip_personal_answers(bank)
    diagnostics = validate_bank(working)
    report.setdefault("files", {})
    report.setdefault("warnings", [])
    word = report["wordExport"] = {"selectedQuestions": len(working["questions"]), "wordQuestions": 0,
        "manualQuestions": 0, "formatManualQuestions": 0, "pendingQuestions": 0,
        "wordQuestionIds": [], "manualQuestionIds": [], "pendingQuestionIds": [],
        "backupManualQuestionIds": [], "needsBriefExplanation": [], "pending": [],
        "formatManual": [], "parts": [], "mediaProblems": [], "mediaAttachments": [],
        "manualTextWarnings": [], "warnings": [], "platformImportVerified": False,
        "qaMode": word_qa_mode, "maxBytes": max_bytes}
    by_question = {}
    for diagnostic in diagnostics:
        by_question.setdefault(diagnostic.get("questionId"), []).append(diagnostic)
    eligible, manual, backup = [], [], []
    media_cache = {}
    word_media = {qid:list(files) for qid, files in media.items()}
    for question in working["questions"]:
        qid, kind = question["id"], _type(question)
        reasons = [i.get("code", "") for i in question.get("issues", []) if i.get("code") in CONTENT_CODES]
        reasons += [i["code"] for i in by_question.get(qid, []) if i["code"] in CONTENT_CODES]
        if kind not in HEADINGS:
            reasons.append("unknown_type")
        if qid in required_missing:
            reasons.append("missing_required_image")
        if not _answers(question):
            reasons.append("missing_correct_answer")
        if not _text(question.get("title")):
            reasons.append("empty_title")
        format_reason = None
        needs_media_backup = False
        try:
            _answer(question)
        except FormatManual as error:
            format_reason = str(error)
        except ValueError as error:
            reasons.append(str(error))
        for problem in _check_media(question, media, out, media_cache):
            word['mediaProblems'].append(dict(problem, questionId=qid))
            if problem['required'] and problem['reason'] == 'unreadable_media':
                reasons.append('invalid_required_image')
            elif problem['reason'] == 'unreadable_media':
                filename = Path(problem['file']).name
                word_media[qid] = [item for item in word_media[qid] if item != filename]
                warning = f"题目 {qid} 的非关键图片 {problem['file']} 无法解码；Word 暂省略该图，原附件及备用人工材料保留。"
                word['warnings'].append(warning)
                report['warnings'].append(warning)
                needs_media_backup = True
            else:
                format_reason = problem['reason']
        try:
            _image_groups(question, media, report)
        except ValueError as error:
            reasons.append(str(error))
        if reasons:
            word["pendingQuestionIds"].append(qid)
            word["pending"].append({"questionId": qid, "reasons": sorted(set(reasons)),
                "sources": question.get("sources", []),
                "suggestions": ["回看列出的原题位置、正确答案及关键媒体；列明可依据的来源与补齐建议，用户同意后再补充并重新核验。"],
                "requiresUserApprovalBeforeSupplement": True})
            continue
        question["_wordBrief"] = _brief(question, report)
        if format_reason:
            manual.append(question)
            word["formatManual"].append({"questionId": qid, "reason": format_reason})
            continue
        if needs_media_backup or media.get(qid):
            backup.append(question)
        if media.get(qid):
            warning = f"题目 {qid} 的图片及位置已保留在导出材料和附件；已测试的 Word 导入曾出现图片占位符，需逐题核对实际显示，备用材料不代表已补图。"
            word['warnings'].append(warning)
            report['warnings'].append(warning)
        if kind == "简答":
            if word_qa_mode == "manual":
                manual.append(question)
            else:
                backup.append(question)
                eligible.append(question)
        else:
            eligible.append(question)
    kinds = list(HEADINGS)
    eligible.sort(key=lambda q: kinds.index(_type(q)))
    report["_wordMedia"] = word_media
    report["files"]["questionDocxParts"] = []
    report["files"].pop("questionDocx", None)
    report["files"].pop("manualQuestions", None)
    report["files"].pop("manualQuestionsMarkdown", None)
    report["wordQuestionOrder"] = []
    try:
        with tempfile.TemporaryDirectory(prefix="word-build-", dir=out) as folder:
            work = Path(folder)
            reference = work / "reference.docx"
            _reference_document(reference)
            from docx import Document
            from docx.shared import Pt
            design = Document(reference)
            for name in ("Normal", "Body Text"):
                style = design.styles[name]
                style.font.size = Pt(11)
                style.paragraph_format.line_spacing = 1.15
                style.paragraph_format.space_after = Pt(3)
            design.styles["Heading 1"].font.size = Pt(14)
            design.save(reference)
            counter = 0
            def split(questions):
                nonlocal counter
                if not questions:
                    return
                counter += 1
                candidate = work / f"candidate-{counter}.docx"
                try:
                    _render_docx(questions, candidate, out, pandoc, reference, report)
                except MissingRenderMedia as error:
                    if len(questions) > 1:
                        midpoint = len(questions) // 2
                        split(questions[:midpoint])
                        split(questions[midpoint:])
                    else:
                        question = questions[0]
                        required = any(item.get('required') and Path(item.get('path', '')).name in error.filenames
                                       for item in question.get('images', []))
                        for filename in error.filenames:
                            word['mediaProblems'].append({'questionId':question['id'], 'file':'assets/'+filename,
                                'reason':'not_embedded_in_docx', 'required':bool(required)})
                        if required:
                            word['pendingQuestionIds'].append(question['id'])
                            word['pending'].append({'questionId':question['id'], 'reasons':['missing_required_image'],
                                'sources':question.get('sources', []),
                                'suggestions':['关键图未出现在生成的 Word 中；回看原图并核对可恢复格式，用户同意后再补齐。'],
                                'requiresUserApprovalBeforeSupplement':True})
                        else:
                            manual.append(question)
                            word['formatManual'].append({'questionId':question['id'], 'reason':'not_embedded_in_docx'})
                    return
                if candidate.stat().st_size > max_bytes:
                    _compact_docx(candidate)
                if candidate.stat().st_size <= max_bytes:
                    part_number = len(word["parts"]) + 1
                    filename = "题库.docx" if part_number == 1 else f"题库-{part_number:02d}.docx"
                    target = out / filename
                    candidate.replace(target)
                    entries = [{"number": n, "questionId": q["id"], "type": _type(q)} for n, q in enumerate(questions, 1)]
                    word["parts"].append({"file": str(target), "bytes": target.stat().st_size, "questions": entries})
                    report["files"]["questionDocxParts"].append(str(target))
                    report["wordQuestionOrder"].extend(dict(entry, file=str(target)) for entry in entries)
                    word["wordQuestionIds"].extend(q["id"] for q in questions)
                elif len(questions) > 1:
                    midpoint = len(questions) // 2
                    split(questions[:midpoint])
                    split(questions[midpoint:])
                else:
                    question = questions[0]
                    manual.append(question)
                    word["formatManual"].append({"questionId": question["id"], "reason": "single_question_exceeds_size_limit",
                                                 "bytes": candidate.stat().st_size})
            split(eligible)
        if word["parts"]:
            report["files"]["questionDocx"] = word["parts"][0]["file"]
        manual_ids = {q["id"] for q in manual}
        backup = list({q['id']:q for q in backup if q["id"] in word['wordQuestionIds'] and q["id"] not in manual_ids}.values())
        word["manualQuestionIds"] = [q["id"] for q in manual]
        word["backupManualQuestionIds"] = [q["id"] for q in backup]
        material = manual + backup
        if material:
            _write_manual_materials(material, out, pandoc, report, media)
        word["wordQuestions"] = len(word["wordQuestionIds"])
        word["manualQuestions"] = len(word["manualQuestionIds"])
        word["pendingQuestions"] = len(word["pendingQuestionIds"])
        word["formatManualQuestions"] = len(word["formatManual"])
        sets = [set(word[key]) for key in ("wordQuestionIds", "manualQuestionIds", "pendingQuestionIds")]
        selected = {q["id"] for q in working["questions"]}
        if set.union(*sets) != selected or any(sets[i] & sets[j] for i in range(3) for j in range(i)):
            raise RuntimeError("Word、人工处理与待核实题目未满足 ID 集合守恒。")
        report["wordFormat"] = "佛脚刷题：题型分组、数字题号、答案与简析"
    finally:
        report.pop("_wordMedia", None)
